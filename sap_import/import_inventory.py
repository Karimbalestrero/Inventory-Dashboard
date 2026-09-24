from pathlib import Path
import shutil
import time
import subprocess
import pandas as pd
from openpyxl import load_workbook
import subprocess

# ============================================================
# PFADE
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

INCOMING_DIR = BASE_DIR / "sap_import" / "incoming"
PROCESSED_DIR = BASE_DIR / "sap_import" / "processed"

TARGET_FILE = BASE_DIR / "data" / "inventory.xlsx"


# ============================================================
# EINSTELLUNGEN
# ============================================================

CHECK_INTERVAL = 3

EXPECTED_COLUMNS = [
    "Phys. Inventory Doc.",
    "Item",
    "Material",
    "Batch",
    "Plant",
    "Storage Location",
    "Physical inventory status",
    "Special Stock",
    "Stock type",
]


# ============================================================
# ORDNER SICHERSTELLEN
# ============================================================

INCOMING_DIR.mkdir(parents=True, exist_ok=True)
PROCESSED_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# PRÜFEN, OB DATEI FERTIG GESCHRIEBEN IST
# ============================================================

def wait_until_file_ready(file_path):

    previous_size = -1

    for _ in range(20):

        try:
            current_size = file_path.stat().st_size

            if current_size > 0 and current_size == previous_size:

                # Zusätzlich testen, ob Excel lesbar ist
                pd.read_excel(file_path, nrows=2)

                return True

            previous_size = current_size

        except Exception:
            pass

        time.sleep(1)

    return False


# ============================================================
# SAP EXPORT VERARBEITEN
# ============================================================
# ============================================================
# INVENTORY.XLSX AUTOMATISCH ZU GITHUB PUSHEN
# ============================================================

def push_inventory_to_github():

    print()
    print("GitHub-Upload wird gestartet...")

    try:
        # Nur inventory.xlsx vormerken
        subprocess.run(
            ["git", "add", "data/inventory.xlsx"],
            cwd=BASE_DIR,
            check=True,
        )

        # Prüfen, ob es überhaupt eine Änderung gibt
        result = subprocess.run(
            ["git", "diff", "--cached", "--quiet"],
            cwd=BASE_DIR,
        )

        if result.returncode == 0:
            print("Keine neuen Änderungen für GitHub.")
            return True

        # Commit erstellen
        subprocess.run(
            [
                "git",
                "commit",
                "-m",
                "Inventurdaten automatisch aktualisiert",
            ],
            cwd=BASE_DIR,
            check=True,
        )

        # Zu GitHub hochladen
        subprocess.run(
            ["git", "push"],
            cwd=BASE_DIR,
            check=True,
        )

        print("GitHub erfolgreich aktualisiert.")
        return True

    except subprocess.CalledProcessError as e:
        print(f"FEHLER beim GitHub-Upload: {e}")
        return False

    except Exception as e:
        print(f"Unerwarteter Fehler beim GitHub-Upload: {e}")
        return False
def process_export(export_file):

    print()
    print("=" * 60)
    print(f"Neue Datei erkannt: {export_file.name}")

    if not wait_until_file_ready(export_file):
        print("Datei konnte noch nicht gelesen werden.")
        return


    # --------------------------------------------------------
    # SAP EXPORT LESEN
    # --------------------------------------------------------

    try:
        sap_df = pd.read_excel(
            export_file,
            dtype=object,
        )

    except Exception as e:
        print(f"Fehler beim Lesen des SAP-Exports: {e}")
        return


    # Spaltennamen bereinigen
    sap_df.columns = (
        sap_df.columns
        .astype(str)
        .str.strip()
    )


    # --------------------------------------------------------
    # SPALTEN PRÜFEN
    # --------------------------------------------------------

    missing = [
        column
        for column in EXPECTED_COLUMNS
        if column not in sap_df.columns
    ]

    if missing:
        print("FEHLER: Im SAP-Export fehlen Spalten:")
        for column in missing:
            print(f"  - {column}")
        return


    # Nur die benötigten 9 Spalten und richtige Reihenfolge
    sap_df = sap_df[EXPECTED_COLUMNS]


    # Komplett leere Zeilen entfernen
    sap_df = sap_df.dropna(how="all")


    print(f"{len(sap_df)} Positionen gefunden.")


    # --------------------------------------------------------
    # BESTEHENDE INVENTORY.XLSX ÖFFNEN
    # --------------------------------------------------------

    if not TARGET_FILE.exists():
        print(f"FEHLER: {TARGET_FILE} wurde nicht gefunden.")
        return


    try:
        workbook = load_workbook(TARGET_FILE)

        worksheet = workbook.active

    except Exception as e:
        print(f"Fehler beim Öffnen von inventory.xlsx: {e}")
        return


    # --------------------------------------------------------
    # ALTE DATEN LÖSCHEN
    #
    # Zeile 1 bleibt erhalten.
    # A2:I... wird geleert.
    # --------------------------------------------------------

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=9,
    ):

        for cell in row:
            cell.value = None


    # --------------------------------------------------------
    # NEUE SAP-DATEN SCHREIBEN
    # --------------------------------------------------------

    for row_number, row in enumerate(
        sap_df.itertuples(index=False, name=None),
        start=2,
    ):

        for column_number, value in enumerate(
            row,
            start=1,
        ):

            # pandas NaN nicht nach Excel schreiben
            if pd.isna(value):
                value = None

            worksheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )


    # --------------------------------------------------------
    # SPEICHERN
    # --------------------------------------------------------

    try:
        workbook.save(TARGET_FILE)

    except PermissionError:

        print()
        print("inventory.xlsx konnte nicht gespeichert werden.")
        print("Bitte prüfen, ob die Datei gerade in Excel geöffnet ist.")

        workbook.close()

        return

    except Exception as e:

        print(f"Fehler beim Speichern: {e}")

        workbook.close()

        return


    workbook.close()


    print(
        f"inventory.xlsx erfolgreich aktualisiert: "
        f"{len(sap_df)} Positionen"
    )


    # --------------------------------------------------------
    # SAP EXPORT ARCHIVIEREN
    # --------------------------------------------------------

    timestamp = time.strftime("%Y%m%d_%H%M%S")

    destination = (
        PROCESSED_DIR
        / f"{export_file.stem}_{timestamp}{export_file.suffix}"
    )

    try:
        shutil.move(
            str(export_file),
            str(destination),
        )

        print(f"Export archiviert: {destination.name}")

    except Exception as e:
        print(f"Export konnte nicht verschoben werden: {e}")
    # --------------------------------------------------------
    # GITHUB AUTOMATISCH AKTUALISIEREN
    # --------------------------------------------------------
   
    github_ok = push_inventory_to_github()
    if github_ok:
        print("Warte auf Aktualisierung des Online-Dashboards...")
        time.sleep(30)

        print("Erstelle aktuellen Dashboard-Screenshot...")
        subprocess.run(
            ["py", "screenshot.py"],
            cwd=str(BASE_DIR),
            check=False,
        )

    if github_ok:
        print("Streamlit-Daten wurden an GitHub übertragen.")
    else:
        print("Excel wurde aktualisiert, aber GitHub-Upload ist fehlgeschlagen.")

# ============================================================
# ORDNER ÜBERWACHEN
# ============================================================

def main():

    print("=" * 60)
    print("SAP Inventory Import")
    print("=" * 60)

    print()
    print(f"Überwache:")
    print(INCOMING_DIR)

    print()
    print("SAP-Export einfach in diesen Ordner speichern.")
    print("Beenden mit STRG+C")
    print()


    while True:

        files = [
    file
    for file in INCOMING_DIR.iterdir()
    if file.is_file()
    and file.suffix.lower() in [".xlsx", ".xlsm"]
    and file.name.upper().startswith("EXPORT_")
    and not file.name.startswith("~$")
]


        if files:

            # Älteste Datei zuerst
            files.sort(
                key=lambda file: file.stat().st_mtime
            )

            for export_file in files:
                process_export(export_file)


        time.sleep(CHECK_INTERVAL)


if __name__ == "__main__":
    main()